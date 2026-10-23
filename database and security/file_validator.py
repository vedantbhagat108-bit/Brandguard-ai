import io
import os
from fastapi import UploadFile, HTTPException, status
import pypdf
import docx

MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024  # 5 MB
ALLOWED_EXTENSIONS = {".txt", ".pdf", ".docx"}


async def process_and_validate_file(file: UploadFile) -> str:
    """
    Asynchronously validates an uploaded file and extracts its readable text.

    Validation rules:
      - File extension whitelist: .txt, .pdf, .docx only. Raises HTTP 400 on unsupported extensions.
      - Maximum file size: 5 MB. Raises HTTP 400 if exceeded.
      - Non-empty check: Raises HTTP 400 if file is 0 bytes.
      - Parsing robustness: Extracts text cleanly using pypdf, python-docx, or UTF-8 decoding.
      - Error handling: Raises HTTP 422 if parsing fails, or HTTP 400 if no readable text is found.

    Returns:
      Cleaned, non-empty extracted string.
    """
    filename = file.filename or ""
    _, ext = os.path.splitext(filename)
    ext = ext.lower()

    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file format '{ext}'. Allowed extensions are: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
        )

    # Read binary payload
    try:
        content: bytes = await file.read()
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Could not read uploaded file: {str(e)}"
        )

    # Validate non-empty payload
    if len(content) == 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Uploaded file is empty (0 bytes)."
        )

    # Validate file size boundary (5 MB)
    if len(content) > MAX_FILE_SIZE_BYTES:
        max_mb = MAX_FILE_SIZE_BYTES // (1024 * 1024)
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"File size exceeds maximum allowed limit of {max_mb} MB (received {len(content) / (1024 * 1024):.2f} MB)."
        )

    # Extract text based on file format
    extracted_text = ""
    try:
        if ext == ".txt":
            try:
                extracted_text = content.decode("utf-8")
            except UnicodeDecodeError:
                extracted_text = content.decode("latin-1", errors="replace")

        elif ext == ".pdf":
            reader = pypdf.PdfReader(io.BytesIO(content))
            pages = []
            for page in reader.pages:
                page_str = page.extract_text()
                if page_str and page_str.strip():
                    pages.append(page_str.strip())
            extracted_text = "\n\n".join(pages)

        elif ext == ".docx":
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_str:
                        paragraphs.append(row_str)
            extracted_text = "\n".join(paragraphs)

    except Exception as parse_err:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Failed to parse document: {str(parse_err)}"
        )

    cleaned = extracted_text.strip()
    if not cleaned:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Uploaded document contains no readable text content."
        )

    return cleaned
